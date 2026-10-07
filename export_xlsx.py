"""Export pendants_flat to a formatted .xlsx, laid out like the original PDF."""
import sqlite3
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

DB_PATH = Path(__file__).parent / "pendants.db"
OUT_PATH = Path(__file__).parent / "Tree of Memory stock - Pendants.xlsx"

COLUMN_WIDTHS = {
    "Parent ID": 11, "SKU": 16, "Photo": 32, "Name": 28, "Description": 50,
    "Price GBP": 10, "Material": 24, "Gemstone": 13, "Cleaning": 40, "Size": 16,
    "Sourcing": 10, "Chain included": 10, "Chain type": 12, "Chain material": 14,
    "Chain length": 14, "Clasp": 10, "Tags": 45,
}


def main():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute("SELECT * FROM pendants_flat")
    columns = [d[0] for d in cur.description]
    rows = cur.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Pendants"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4A7C59", end_color="4A7C59", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(name, 16)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap
        ws.row_dimensions[row_idx].height = 90

    ws.freeze_panes = "A2"
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({len(rows)} rows)")


if __name__ == "__main__":
    main()
